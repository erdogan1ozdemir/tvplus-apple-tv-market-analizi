#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Rakip Trafiği Excel çıktısı.

Panoda gösterilenin tamamı + panoya sığmayan ham SERP tablosu tek dosyada.
Kullanım: python3 scripts/excel_rakip.py
"""
import json, os, collections, sys
sys.path.insert(0, os.path.dirname(__file__))
from openpyxl import Workbook
from excel_ortak import sayfa_yaz, not_sayfasi

serp = json.load(open("data/serp/serp.json", encoding="utf-8"))
IZLEME = {"Korsan", "Meşru Platform"}

# ——— dizi bazında
d = {}
for s in serp:
    o = d.setdefault(s["dizi"], dict(dizi=s["dizi"], kw=s["keyword"], tur=s["tur"],
        kategori=s["kategori"], hacim=0, izleme=0, korsan=0, mesru=0, agregat=0,
        ilk3=0, jenerik=0, sayfa=0))
    o["hacim"] = max(o["hacim"], s["aylik_hacim"] or 0)
    o["sayfa"] += 1
    t = s["sayfa_trafik"] or 0
    if not s["dizi_sayfasi"]:
        o["jenerik"] += t; continue
    if s["sinif"] in IZLEME: o["izleme"] += t
    if s["sinif"] == "Korsan": o["korsan"] += t
    elif s["sinif"] == "Meşru Platform": o["mesru"] += t
    else: o["agregat"] += t
    if s["pozisyon"] <= 3: o["ilk3"] += t
dizi = sorted(d.values(), key=lambda x: -x["izleme"])

# ——— domain bazında
dm = {}
for s in serp:
    o = dm.setdefault(s["domain"], dict(dom=s["domain"], sinif=s["sinif"], n=0, ilk3=0,
                                        trafik=0, diziler=set()))
    o["n"] += 1
    if s["pozisyon"] <= 3: o["ilk3"] += 1
    o["diziler"].add(s["dizi"])
    if s["dizi_sayfasi"]: o["trafik"] += s["sayfa_trafik"] or 0
domain = sorted(dm.values(), key=lambda x: -x["n"])

wb = Workbook(); wb.remove(wb.active)

SAY = {"Sorgu Hacmi","İzleme Odaklı Trafik","Korsan","Meşru Platform","Agregatör",
       "İlk 3 Trafiği","Taranan Sayfa","Pozisyon",
       "Ahrefs Sayfa Trafiği","İlk 10'da","İlk 3'te","Dizi Sayısı","Aylık Hacim"}

sayfa_yaz(wb.create_sheet("Dizi Bazında Trafik"),
    ["Dizi","Sorgu","Kategori","Tür","Sorgu Hacmi","İzleme Odaklı Trafik","Korsan",
     "Meşru Platform","Agregatör","İlk 3 Trafiği","Taranan Sayfa"],
    [[x["dizi"], x["kw"], x["kategori"], x["tur"], x["hacim"], x["izleme"], x["korsan"],
      x["mesru"], x["agregat"], x["ilk3"], x["sayfa"]] for x in dizi], SAY)

sayfa_yaz(wb.create_sheet("Tüm SERP Sonuçları"),
    ["Dizi","Sorgu","Aylık Hacim","Pozisyon","Domain","Sayfa URL","Sayfa Başlığı",
     "Sınıf","Ahrefs Sayfa Trafiği","Kapsam"],
    [[s["dizi"], s["keyword"], s["aylik_hacim"], s["pozisyon"], s["domain"], s["url"],
      s.get("baslik") or "", s["sinif"], s["sayfa_trafik"] or 0,
      "Dizi sayfası" if s["dizi_sayfasi"] else "Jenerik / başka yapım"] for s in serp], SAY,
    {"Sayfa URL": 60, "Sayfa Başlığı": 45})

sayfa_yaz(wb.create_sheet("Domain Özeti"),
    ["Domain","Sınıf","İlk 10'da","İlk 3'te","Dizi Sayısı","Ahrefs Sayfa Trafiği"],
    [[x["dom"], x["sinif"], x["n"], x["ilk3"], len(x["diziler"]), x["trafik"]] for x in domain], SAY)

sinif = collections.Counter(); sinifT = collections.Counter()
for s in serp:
    sinif[s["sinif"]] += 1
    if s["dizi_sayfasi"]: sinifT[s["sinif"]] += s["sayfa_trafik"] or 0
sayfa_yaz(wb.create_sheet("Sınıf Özeti"),
    ["Sınıf","İlk 10'da","Ahrefs Sayfa Trafiği"],
    [[k, v, sinifT[k]] for k, v in sinif.most_common()], SAY)

not_sayfasi(wb.create_sheet("Yöntem"), [
 ("Kapsam", f"{len({s['dizi'] for s in serp})} dizi · {len({s['keyword'] for s in serp})} sorgu · "
            f"{len(serp)} SERP sonucu. Her dizi için '{{dizi adı}} izle' sorgusunun ilk 10 organik sonucu."),
 ("Sıralama kaynağı", "DataForSEO · Google Türkiye (location 2792, dil tr, masaüstü), ilk 10 organik sonuç."),
 ("Trafik kaynağı", "Ahrefs batch-analysis, mode=exact. Değer o URL'in tahmini aylık organik trafiğidir; "
                    "yalnızca bu sorgudan geleni değil, sayfanın tamamını kapsar. Üst sınır okuması verir."),
 ("İzleme odaklı trafik", "Korsan + meşru platform sayfalarının toplamı. Agregatör ve bilgi siteleri "
                          "(JustWatch, IMDb, Wikipedia) izleme sayfası sunmadığı için ayrı sütundadır."),
 ("Kapsam dışı sonuçlar", "Aramanın ilk 10'una dizinin kendi sayfası olmayan sonuçlar da giriyor: "
                   "platformların genel katalog sayfaları (nowtv.com.tr/film-izle gibi, kendi bütün "
                   "trafiğini getirir) ve aynı adı taşıyan başka yapımlar. Bu satırların trafiği dizi "
                   "toplamlarına katılmaz; 'Tüm SERP Sonuçları' sayfasında Kapsam sütununda "
                   "'Jenerik / başka yapım' olarak işaretlidir."),
 ("Sınır", "Ad belirsizliği taşıyan dizilerde sayfa düzeyinde sızıntı sürebilir; ör. 'now and then' "
           "sorgusunda 1995 tarihli aynı adlı filmin platform sayfası ilk 10'a giriyor."),
 ("Kaynak", "Kaynak: DataForSEO (sıralama) · Ahrefs (sayfa trafiği)"),
])

os.makedirs("cikti", exist_ok=True)
yol = "cikti/appletv-rakip-trafigi.xlsx"
wb.save(yol)
print(f"{yol} · {len(dizi)} dizi · {len(serp)} SERP satırı · {len(domain)} domain")
