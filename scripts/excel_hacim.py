#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Arama hacmi Excel çıktısı: master liste + küme bazında hacim sayfaları.

Kullanım: python3 scripts/excel_hacim.py
"""
import csv, glob, os, sys, collections
sys.path.insert(0, os.path.dirname(__file__))
from openpyxl import Workbook
from excel_ortak import sayfa_yaz, not_sayfasi

satir = []
for y in sorted(glob.glob("data/raw/hacim_*.csv")):
    satir += list(csv.DictReader(open(y, encoding="utf-8-sig")))
veri    = [r for r in satir if r.get("veri_var") == "evet"]
kapsam  = [r for r in veri  if not r.get("mantik_denetim")]
disi    = [r for r in veri  if r.get("mantik_denetim")]
V = lambda r: int(float(r.get("search_volume") or 0))

AYLAR = [k for k in satir[0] if len(k) == 7 and k[4] == "-"]
AYLAR.sort()
SON12, ONC12 = AYLAR[-12:], AYLAR[-24:-12]
def top(r, aylar): return sum(int(float(r.get(a) or 0)) for a in aylar)

SAY = {"Aylık Hacim","Son 12 Ay","Önceki 12 Ay","Son 12 Ay Aylık Ort.","YoY %",
       "Keyword","Dizi Sayısı","Toplam Aylık Hacim","İzleme Hacmi","İzleme Payı %",
       "Çıkış Yılı","Sezon Sayısı"}

wb = Workbook(); wb.remove(wb.active)

# ——— Master liste
sayfa_yaz(wb.create_sheet("Master Liste"),
  ["Keyword","Dizi","Kategori","Tür","Alt Tür","Sayfa Tipi","Intent","Varlık Tipi",
   "Marka Tipi","Dil","Çıkış Yılı","Sezon Sayısı","Durum",
   "Aylık Hacim","Son 12 Ay","Önceki 12 Ay","YoY %"],
  [[r["keyword"], r["dizi"], r["kategori"], r["tur"], r.get("tur_ham",""), r["sayfa_tipi"],
    r["intent_katmani"], r["entity_tipi"], r["marka_tipi"], r["dil"],
    r["cikis_yili"], r["sezon_sayisi"], r["durum"], V(r), top(r, SON12), top(r, ONC12),
    round(100*(top(r,SON12)-top(r,ONC12))/top(r,ONC12), 1) if top(r, ONC12) else ""]
   for r in sorted(kapsam, key=lambda x: -V(x))], SAY, {"Keyword": 42})

# ——— Küme sayfaları
def kume(ad, alan, ek=None):
    g = collections.defaultdict(lambda: {"kw":0, "hac":0, "izl":0, "dizi":set()})
    for r in kapsam:
        o = g[r[alan] or "–"]; o["kw"] += 1; o["hac"] += V(r); o["dizi"].add(r["dizi"])
        if r["intent_katmani"] == "İzleme": o["izl"] += V(r)
    sat = [[k, v["kw"], len(v["dizi"]), v["hac"], v["izl"],
            round(100*v["izl"]/v["hac"], 1) if v["hac"] else 0]
           for k, v in sorted(g.items(), key=lambda x: -x[1]["hac"])]
    sayfa_yaz(wb.create_sheet(ad),
      [ek or ad, "Keyword", "Dizi Sayısı", "Toplam Aylık Hacim", "İzleme Hacmi", "İzleme Payı %"],
      sat, SAY | {"İzleme Payı %"})

kume("Dizi", "dizi"); kume("Kategori", "kategori"); kume("Tür", "tur")
kume("Sayfa Tipi", "sayfa_tipi"); kume("Intent", "intent_katmani")
kume("Varlık Tipi", "entity_tipi"); kume("Marka Tipi", "marka_tipi")

# ——— Dizi × sayfa tipi çapraz tablosu
tipler = sorted({r["sayfa_tipi"] for r in kapsam})
capraz = collections.defaultdict(lambda: collections.Counter())
for r in kapsam: capraz[r["dizi"]][r["sayfa_tipi"]] += V(r)
sayfa_yaz(wb.create_sheet("Dizi × Sayfa Tipi"),
  ["Dizi", "Toplam Aylık Hacim"] + tipler,
  [[d, sum(c.values())] + [c.get(t, 0) for t in tipler]
   for d, c in sorted(capraz.items(), key=lambda x: -sum(x[1].values()))],
  SAY | set(tipler))

# ——— Kapsam dışı
sayfa_yaz(wb.create_sheet("Kapsam Dışı"),
  ["Keyword","Dizi","Aylık Hacim","Gerekçe"],
  [[r["keyword"], r["dizi"], V(r), r["mantik_denetim"]]
   for r in sorted(disi, key=lambda x: -V(x))], SAY, {"Gerekçe": 70, "Keyword": 36})

not_sayfasi(wb.create_sheet("Yöntem"), [
 ("Kapsam", f"{len(kapsam):,} keyword · {len({r['dizi'] for r in kapsam})} dizi · "
            f"aylık toplam {sum(map(V, kapsam)):,} arama.".replace(",", ".")),
 ("Dönem", f"{AYLAR[0]} – {AYLAR[-1]} ({len(AYLAR)} ay). Son 12 Ay = {SON12[0]}–{SON12[-1]}, "
           f"Önceki 12 Ay = {ONC12[0]}–{ONC12[-1]}. YoY bu iki pencere arasındadır."),
 ("Aylık Hacim", "Google Ads'in verdiği aylık ortalama arama hacmi; aylık seri sütunlarının "
                 "ortalaması değil, aracın kendi değeridir."),
 ("Kapsam dışı", f"{len(disi)} keyword mantık denetimiyle kapsam dışı bırakıldı "
                 "(ad belirsizliği ya da dizi olmayan kayıt). Satırlar silinmedi, "
                 "'Kapsam Dışı' sayfasında gerekçesiyle durur."),
 ("İzleme Payı", "İzleme intent'indeki hacmin küme toplamına oranı."),
 ("Kaynak", "Kaynak: DataForSEO Google Ads Search Volume · Türkiye · Türkçe"),
])

os.makedirs("cikti", exist_ok=True)
yol = "cikti/appletv-arama-hacmi.xlsx"
wb.save(yol)
print(f"{yol} · {len(kapsam)} keyword · {len(wb.sheetnames)} sayfa: {', '.join(wb.sheetnames)}")
