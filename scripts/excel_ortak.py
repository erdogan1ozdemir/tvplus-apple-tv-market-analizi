#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel çıktıları için ortak biçim.

Inbound Design System · Bölüm 15.3 (Excel çalışma aracı rejimi):
başlık satırı Shadowed Charcoal #434343 + beyaz kalın, gövde Calibri + ink teal
#10332F, dolgusuz delta, dikeyde ortalı hücreler, satır başına üst kenarlık.
Office varsayılan teması (lacivert başlık, sarı highlight) kullanılmaz.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASLIK_DOLGU = PatternFill("solid", fgColor="434343")
BASLIK_YAZI  = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
GOVDE_YAZI   = Font(name="Calibri", size=11, color="10332F")
NOT_YAZI     = Font(name="Calibri", size=11, bold=True, color="FF7B52")
YESIL        = Font(name="Calibri", size=11, bold=True, color="2E7D32")
KIRMIZI      = Font(name="Calibri", size=11, bold=True, color="D32F2F")
UST_KENAR    = Border(top=Side(style="thin", color="E0E0E0"))
ORTA_SOL     = Alignment(horizontal="left",   vertical="center", wrap_text=False)
ORTA_ORTA    = Alignment(horizontal="center", vertical="center")

def sayfa_yaz(ws, basliklar, satirlar, sayisal=(), genislik=None, dondur=True):
    """Başlık + gövde yazar ve ev biçimini uygular. `sayisal` sütun adları ortalanır."""
    ws.append(basliklar)
    for h in ws[1]:
        h.fill = BASLIK_DOLGU; h.font = BASLIK_YAZI; h.alignment = ORTA_ORTA
    say_idx = {i for i, b in enumerate(basliklar) if b in sayisal}
    for s in satirlar:
        ws.append(list(s))
    for satir in ws.iter_rows(min_row=2):
        for i, h in enumerate(satir):
            h.font = GOVDE_YAZI
            h.alignment = ORTA_ORTA if i in say_idx else ORTA_SOL
            h.border = UST_KENAR
            if isinstance(h.value, (int, float)) and i in say_idx:
                h.number_format = "#,##0"
    ws.freeze_panes = "A2" if dondur else None
    for i, b in enumerate(basliklar):
        w = (genislik or {}).get(b)
        if not w:
            en = max([len(str(b))] + [len(str(s[i])) for s in satirlar[:400]] or [8])
            w = min(max(en + 2, 10), 60)
        ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = w
    ws.auto_filter.ref = ws.dimensions
    return ws

def not_sayfasi(ws, satirlar):
    """Yöntem/açıklama sayfası: etiket kalın coral, gövde ink teal."""
    for etiket, metin in satirlar:
        ws.append([etiket, metin])
    for satir in ws.iter_rows():
        satir[0].font = NOT_YAZI
        satir[0].alignment = Alignment(vertical="top", horizontal="left")
        satir[1].font = GOVDE_YAZI
        satir[1].alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    return ws
